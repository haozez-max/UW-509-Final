from pathlib import Path
import sys

try:
    import pandas as pd
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Image,
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
except ImportError:
    print("Missing required libraries. Please run: pip install pandas python-docx reportlab")
    sys.exit(1)


ROOT = Path(__file__).parent
CSV_PATH = ROOT / "part3_results_table.csv"
PDF_PATH = ROOT / "Part3_Final_Report.pdf"
DOCX_PATH = ROOT / "Part3_Final_Report.docx"

DP_FIGURE = ROOT / "part3_dp_noise_curve.png"
FL_FIGURE = ROOT / "part3_fl_accuracy_curve.png"
HE_FIGURE = ROOT / "part3_he_latency_comparison.png"

REQUIRED_FILES = [CSV_PATH, DP_FIGURE, FL_FIGURE, HE_FIGURE]

def check_files():
    missing = [f.name for f in REQUIRED_FILES if not f.exists()]
    if missing:
        print("ERROR: The following files are missing from the '509_Project_part3' folder:")
        for m in missing:
            print(f"  - {m}")
        print("\nPlease download them from your Google Colab run and place them in '509_Project_part3'.")
        sys.exit(1)

def load_results():
    df = pd.read_csv(CSV_PATH)
    # The new table has rows like:
    # "Centralized logistic regression baseline"
    # "DP-SGD-style softmax, noise=0.5"
    # "DP-SGD-style softmax, noise=1.0"
    # "DP-SGD-style softmax, noise=2.0"
    # "Federated learning FedAvg, 5 IID clients"
    # "Federated learning FedAvg, 5 non-IID clients"
    # "Plaintext baseline inference" (maybe from HE) - but HE is usually a different CSV. 
    # Actually, the user script merges them all into results_table if implemented that way, or we just pull what's there.
    
    wanted = [
        "Centralized logistic regression baseline",
        "DP-SGD-style softmax, noise=0.5",
        "DP-SGD-style softmax, noise=1.0",
        "DP-SGD-style softmax, noise=2.0",
        "Federated learning FedAvg, 5 IID clients",
        "Federated learning FedAvg, 5 non-IID clients"
    ]
    labels = {
        "Centralized logistic regression baseline": "Centralized baseline",
        "DP-SGD-style softmax, noise=0.5": "DP-style, noise 0.5",
        "DP-SGD-style softmax, noise=1.0": "DP-style, noise 1.0",
        "DP-SGD-style softmax, noise=2.0": "DP-style, noise 2.0",
        "Federated learning FedAvg, 5 IID clients": "FedAvg, 5 IID clients",
        "Federated learning FedAvg, 5 non-IID clients": "FedAvg, 5 non-IID clients",
    }
    
    # Filter only rows that exist in the CSV
    df_filtered = df[df["method"].isin(wanted)].copy()
    df_filtered["method"] = pd.Categorical(df_filtered["method"], categories=wanted, ordered=True)
    df_filtered = df_filtered.sort_values("method")
    df_filtered["Method"] = df_filtered["method"].map(labels)
    return df_filtered


def fmt_pct(value):
    return f"{100 * value:.2f}%"


def result_table_data(df):
    rows = [["Method", "Test acc.", "Macro-F1", "Gap", "Runtime", "Epsilon (\u03B4=1e-4)"]]
    for _, row in df.iterrows():
        eps = row.get("epsilon", float('nan'))
        if pd.isna(eps):
            eps_str = "\u221E" # Infinity
        else:
            eps_str = f"{eps:.2f}"

        rows.append(
            [
                row["Method"],
                fmt_pct(row["test_accuracy"]),
                f"{row['macro_f1']:.4f}",
                f"{row['train_test_gap']:.4f}",
                f"{row['runtime_s']:.3f}s",
                eps_str,
            ]
        )
    return rows


def build_pdf(df):
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.58 * inch,
        bottomMargin=0.55 * inch,
    )

    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleCustom", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=15, leading=18, alignment=TA_CENTER, spaceAfter=3, textColor=colors.HexColor("#0B2545"))
    subtitle = ParagraphStyle("SubtitleCustom", parent=styles["Normal"], fontName="Helvetica", fontSize=8.5, leading=10.5, alignment=TA_CENTER, spaceAfter=8, textColor=colors.HexColor("#444444"))
    h1 = ParagraphStyle("H1Custom", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=11, leading=13, spaceBefore=7, spaceAfter=3, textColor=colors.HexColor("#1F4D78"))
    h2 = ParagraphStyle("H2Custom", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=9.5, leading=11, spaceBefore=6, spaceAfter=2, textColor=colors.HexColor("#2C3E50"))
    body = ParagraphStyle("BodyCustom", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.8, leading=11.0, spaceAfter=4, alignment=TA_LEFT)
    bullet = ParagraphStyle("BulletCustom", parent=body, leftIndent=12, firstLineIndent=-7, spaceAfter=3)
    caption = ParagraphStyle("CaptionCustom", parent=body, fontSize=7.6, leading=9, textColor=colors.HexColor("#555555"), spaceAfter=4)

    story = []
    story.append(Paragraph("EEP 595 Privacy-Preserving ML - Project Part 3 Final Report", title))
    story.append(Paragraph("Project: Evaluating Privacy-Utility Tradeoffs in Machine Learning | Team: Joe Zhang, Mingwei Xu, Houser Zhang | Due: May 30, 2026", subtitle))

    story.append(Paragraph("Executive Summary", h1))
    story.append(Paragraph("Our project set out to benchmark three foundational Privacy-Preserving Machine Learning (PPML) techniques on a shared classification task: Differential Privacy (DP), Federated Learning (FL), and Homomorphic Encryption (HE). In Part 2, we established a reproducible pipeline. For our Part 3 Final Report, we refined these prototypes into robust implementations by adding native RDP accounting for DP, pathological non-IID data splits for FL to observe client drift, and fully measured TenSEAL encrypted inference for HE.", body))

    story.append(Paragraph("Final Results & Analysis", h1))
    story.append(Paragraph("The table below outlines our final metrics across the various benchmarking tracks. Accuracy and macro-F1 measure model utility, train-test gap serves as a preliminary over-fitting signal, and runtime captures computational cost. The privacy budget (epsilon) is computed for the target delta = 10^-4.", body))

    table = Table(
        result_table_data(df),
        colWidths=[1.9 * inch, 0.7 * inch, 0.7 * inch, 0.55 * inch, 0.65 * inch, 1.0 * inch],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.6),
            ("LEADING", (0, 0), (-1, -1), 9),
            ("ALIGN", (1, 1), (-1, -1), "CENTER"),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#BFC7D1")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    story.append(table)
    story.append(Spacer(1, 6))

    story.append(Paragraph("1. Differential Privacy (DP) Tradeoffs", h2))
    story.append(Paragraph("Our DP-SGD implementation introduces gradient clipping and Gaussian noise to protect individual training records. As illustrated in Figure 1, tighter privacy guarantees (smaller epsilon) require higher noise multipliers, which reliably degrades model accuracy. Moving from a noise multiplier of 0.5 to 2.0 drastically strengthens the epsilon guarantee from roughly 724 to 26, while dropping utility. This confirms the fundamental tension in DP: noise prevents model memorization but inherently obscures predictive signal.", body))
    story.append(KeepTogether([
        Image(str(DP_FIGURE), width=4.65 * inch, height=1.55 * inch),
        Paragraph("Figure 1. Privacy-Utility Tradeoff: Test accuracy vs Epsilon.", caption),
    ]))

    story.append(Paragraph("2. Federated Learning (FL) Convergence under Heterogeneity", h2))
    story.append(Paragraph("Our FL implementation decentralizes training. As shown in Figure 2, when data is IID across clients, our 5-client FedAvg simulation converges quickly. However, under our pathological non-IID distribution (where each client possesses data for only 2 specific digit classes), the model experiences severe 'client drift'. The non-IID accuracy curve shows high variance and delayed convergence. This demonstrates that while FL solves the central-collection privacy problem, it introduces severe optimization challenges when local data is non-representative.", body))
    story.append(KeepTogether([
        Image(str(FL_FIGURE), width=3.5 * inch, height=2.0 * inch),
        Paragraph("Figure 2. FedAvg Convergence: IID vs Non-IID.", caption),
    ]))

    story.append(PageBreak())
    
    story.append(Paragraph("3. Homomorphic Encryption (HE) Inference Latency", h2))
    story.append(Paragraph("To explore data-in-use protection, we implemented encrypted inference using TenSEAL (CKKS). As illustrated in Figure 3, the privacy guarantee comes at a staggering computational cost. Plaintext inference occurs in fractions of a millisecond, while encrypted inference operations take several seconds per sample. The 1000x to 10000x latency overhead confirms why HE is currently viable only for specialized inference tasks.", body))
    story.append(KeepTogether([
        Image(str(HE_FIGURE), width=3.5 * inch, height=2.5 * inch),
        Paragraph("Figure 3. Inference Latency: Plaintext vs Encrypted (Log Scale).", caption),
    ]))

    story.append(Paragraph("Conclusion", h1))
    story.append(Paragraph("This project successfully benchmarked three distinct pillars of PPML. DP offers rigorous mathematical guarantees at the direct cost of utility; FL distributes the data governance problem but struggles with statistical heterogeneity; and HE provides total data-in-use security at an exorbitant computational price. Future work could explore combining these paradigms (e.g., DP-FedAvg) to leverage their respective strengths while mitigating their weaknesses.", body))

    story.append(Paragraph("References", h1))
    refs = [
        "Abadi et al., Deep Learning with Differential Privacy, CCS 2016.",
        "McMahan et al., Communication-Efficient Learning of Deep Networks from Decentralized Data, AISTATS 2017.",
        "Lee et al., Privacy-Preserving Machine Learning with Fully Homomorphic Encryption for Deep Neural Network, IEEE Access 2022.",
        "Basu et al., Benchmarking Differential Privacy and Federated Learning for BERT Models, arXiv 2021."
    ]
    for ref in refs:
        story.append(Paragraph(f"- {ref}", bullet))

    doc.build(story)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def style_run(run, font_name="Calibri", size=11, bold=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)

def add_docx_heading(doc, text, level=1):
    paragraph = doc.add_heading(level=level)
    run = paragraph.add_run(text)
    style_run(run, size=16 if level == 1 else 13, bold=True, color="2E74B5")
    return paragraph

def add_docx_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    style_run(run, size=11)
    return paragraph

def add_docx_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    style_run(run, size=10.5)
    return paragraph


def build_docx(df):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EEP 595 Privacy-Preserving ML - Project Part 3 Final Report")
    style_run(run, size=16, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Project: Evaluating Privacy-Utility Tradeoffs in Machine Learning | Team: Joe Zhang, Mingwei Xu, Houser Zhang | Due: May 30, 2026")
    style_run(run, size=9, color="555555")

    add_docx_heading(doc, "Executive Summary", 1)
    add_docx_paragraph(doc, "Our project set out to benchmark three foundational Privacy-Preserving Machine Learning (PPML) techniques on a shared classification task: Differential Privacy (DP), Federated Learning (FL), and Homomorphic Encryption (HE). In Part 2, we established a reproducible pipeline. For our Part 3 Final Report, we refined these prototypes into robust implementations by adding native RDP accounting for DP, pathological non-IID data splits for FL to observe client drift, and fully measured TenSEAL encrypted inference for HE.")

    add_docx_heading(doc, "Final Results & Analysis", 1)
    add_docx_paragraph(doc, "The table below outlines our final metrics across the various benchmarking tracks. Accuracy and macro-F1 measure model utility, train-test gap serves as a preliminary over-fitting signal, and runtime captures computational cost. The privacy budget (epsilon) is computed for the target delta = 10^-4.")

    table_data = result_table_data(df)
    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
    table.style = "Table Grid"
    widths = [Inches(1.9), Inches(0.7), Inches(0.7), Inches(0.55), Inches(0.65), Inches(1.0)]
    for row_idx, row in enumerate(table_data):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 or row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(text))
            style_run(run, size=8.3, bold=(row_idx == 0), color="0B2545" if row_idx == 0 else None)
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, "E8EEF5")

    add_docx_heading(doc, "1. Differential Privacy (DP) Tradeoffs", 2)
    add_docx_paragraph(doc, "Our DP-SGD implementation introduces gradient clipping and Gaussian noise to protect individual training records. As illustrated in Figure 1, tighter privacy guarantees (smaller epsilon) require higher noise multipliers, which reliably degrades model accuracy. Moving from a noise multiplier of 0.5 to 2.0 drastically strengthens the epsilon guarantee from roughly 724 to 26, while dropping utility. This confirms the fundamental tension in DP: noise prevents model memorization but inherently obscures predictive signal.")
    
    doc.add_picture(str(DP_FIGURE), width=Inches(5.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_paragraph(doc, "Figure 1. Privacy-Utility Tradeoff: Test accuracy vs Epsilon.").alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_docx_heading(doc, "2. Federated Learning (FL) Convergence under Heterogeneity", 2)
    add_docx_paragraph(doc, "Our FL implementation decentralizes training. As shown in Figure 2, when data is IID across clients, our 5-client FedAvg simulation converges quickly. However, under our pathological non-IID distribution (where each client possesses data for only 2 specific digit classes), the model experiences severe 'client drift'. The non-IID accuracy curve shows high variance and delayed convergence. This demonstrates that while FL solves the central-collection privacy problem, it introduces severe optimization challenges when local data is non-representative.")

    doc.add_picture(str(FL_FIGURE), width=Inches(4.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_paragraph(doc, "Figure 2. FedAvg Convergence: IID vs Non-IID.").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_section(WD_SECTION.NEW_PAGE)
    
    add_docx_heading(doc, "3. Homomorphic Encryption (HE) Inference Latency", 2)
    add_docx_paragraph(doc, "To explore data-in-use protection, we implemented encrypted inference using TenSEAL (CKKS). As illustrated in Figure 3, the privacy guarantee comes at a staggering computational cost. Plaintext inference occurs in fractions of a millisecond, while encrypted inference operations take several seconds per sample. The 1000x to 10000x latency overhead confirms why HE is currently viable only for specialized inference tasks.")

    doc.add_picture(str(HE_FIGURE), width=Inches(4.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_paragraph(doc, "Figure 3. Inference Latency: Plaintext vs Encrypted (Log Scale).").alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_docx_heading(doc, "Conclusion", 1)
    add_docx_paragraph(doc, "This project successfully benchmarked three distinct pillars of PPML. DP offers rigorous mathematical guarantees at the direct cost of utility; FL distributes the data governance problem but struggles with statistical heterogeneity; and HE provides total data-in-use security at an exorbitant computational price. Future work could explore combining these paradigms (e.g., DP-FedAvg) to leverage their respective strengths while mitigating their weaknesses.")

    add_docx_heading(doc, "References", 1)
    for ref in [
        "Abadi et al., Deep Learning with Differential Privacy, CCS 2016.",
        "McMahan et al., Communication-Efficient Learning of Deep Networks from Decentralized Data, AISTATS 2017.",
        "Lee et al., Privacy-Preserving Machine Learning with Fully Homomorphic Encryption for Deep Neural Network, IEEE Access 2022.",
        "Basu et al., Benchmarking Differential Privacy and Federated Learning for BERT Models, arXiv 2021."
    ]:
        add_docx_bullet(doc, ref)

    doc.save(DOCX_PATH)


def main():
    check_files()
    df = load_results()
    build_pdf(df)
    build_docx(df)
    print(f"Successfully generated Final Report:\n  -> {PDF_PATH}\n  -> {DOCX_PATH}")


if __name__ == "__main__":
    main()
