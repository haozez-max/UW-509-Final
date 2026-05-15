from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(".")
CSV_PATH = ROOT / "part2_results_table.csv"
PDF_PATH = ROOT / "Part2_Progress_Report.pdf"
DOCX_PATH = ROOT / "Part2_Progress_Report.docx"
DP_FIGURE = ROOT / "part2_dp_noise_curve.png"
FL_FIGURE = ROOT / "part2_fl_accuracy_curve.png"


def load_results():
    df = pd.read_csv(CSV_PATH)
    wanted = [
        "Centralized logistic regression baseline",
        "DP-SGD-style softmax, noise=0.5",
        "DP-SGD-style softmax, noise=1.0",
        "DP-SGD-style softmax, noise=2.0",
        "Federated learning FedAvg, 5 IID clients",
    ]
    labels = {
        "Centralized logistic regression baseline": "Centralized baseline",
        "DP-SGD-style softmax, noise=0.5": "DP-style, noise 0.5",
        "DP-SGD-style softmax, noise=1.0": "DP-style, noise 1.0",
        "DP-SGD-style softmax, noise=2.0": "DP-style, noise 2.0",
        "Federated learning FedAvg, 5 IID clients": "FedAvg, 5 IID clients",
    }
    df = df[df["method"].isin(wanted)].copy()
    df["method"] = pd.Categorical(df["method"], categories=wanted, ordered=True)
    df = df.sort_values("method")
    df["Method"] = df["method"].map(labels)
    return df


def fmt_pct(value):
    return f"{100 * value:.2f}%"


def result_table_data(df):
    rows = [["Method", "Test acc.", "Macro-F1", "Gap", "MIA AUC", "Runtime"]]
    for _, row in df.iterrows():
        rows.append(
            [
                row["Method"],
                fmt_pct(row["test_accuracy"]),
                f"{row['macro_f1']:.4f}",
                f"{row['train_test_gap']:.4f}",
                f"{row['mia_auc_confidence']:.4f}",
                f"{row['runtime_s']:.3f}s",
            ]
        )
    return rows


def build_pdf(df):
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.58 * inch,
        bottomMargin=0.55 * inch,
    )

    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "TitleCustom",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        alignment=TA_CENTER,
        spaceAfter=3,
        textColor=colors.HexColor("#0B2545"),
    )
    subtitle = ParagraphStyle(
        "SubtitleCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=10.5,
        alignment=TA_CENTER,
        spaceAfter=8,
        textColor=colors.HexColor("#444444"),
    )
    h1 = ParagraphStyle(
        "H1Custom",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=13,
        spaceBefore=7,
        spaceAfter=3,
        textColor=colors.HexColor("#1F4D78"),
    )
    body = ParagraphStyle(
        "BodyCustom",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.8,
        leading=11.0,
        spaceAfter=4,
        alignment=TA_LEFT,
    )
    bullet = ParagraphStyle(
        "BulletCustom",
        parent=body,
        leftIndent=12,
        firstLineIndent=-7,
        spaceAfter=3,
    )
    caption = ParagraphStyle(
        "CaptionCustom",
        parent=body,
        fontSize=7.6,
        leading=9,
        textColor=colors.HexColor("#555555"),
        spaceAfter=4,
    )

    story = []
    story.append(Paragraph("EEP 595 Privacy-Preserving ML - Project Part 2 Progress Report", title))
    story.append(
        Paragraph(
            "Project: Evaluating Privacy-Utility Tradeoffs in Machine Learning | "
            "Team: Joe Zhang, Mingwei Xu, Houser Zhang | Due: May 16, 2026",
            subtitle,
        )
    )

    story.append(Paragraph("Progress and Execution", h1))
    story.append(
        Paragraph(
            "Our proposal committed to benchmarking differential privacy (DP), federated learning (FL), "
            "and homomorphic encryption (HE) on a shared classification task. For Part 2, we built a "
            "runnable Google Colab pipeline that validates the experimental structure before scaling to "
            "larger models and datasets. The current implementation uses scikit-learn's digits dataset "
            "(1,797 grayscale 8x8 handwritten digit images), normalizes pixels to [0, 1], and uses an "
            "80/20 stratified train-test split with seed 509.",
            body,
        )
    )
    story.append(
        Paragraph(
            "Completed components include: a non-private centralized logistic-regression baseline; a "
            "NumPy softmax classifier with per-example gradient clipping and Gaussian noise as a "
            "DP-SGD-style prototype; a five-client IID FedAvg simulation with 30 communication rounds "
            "and two local epochs per round; and an HE feasibility track focused on encrypted inference "
            "rather than encrypted training.",
            body,
        )
    )

    story.append(Paragraph("Preliminary Results", h1))
    story.append(
        Paragraph(
            "The table below is intended for direct use in the final report. Accuracy and macro-F1 measure "
            "utility, train-test gap and confidence-based membership-inference AUC provide preliminary "
            "privacy-leakage proxies, and runtime records practical cost.",
            body,
        )
    )

    table = Table(
        result_table_data(df),
        colWidths=[1.95 * inch, 0.73 * inch, 0.70 * inch, 0.55 * inch, 0.63 * inch, 0.63 * inch],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.6),
                ("LEADING", (0, 0), (-1, -1), 9),
                ("ALIGN", (1, 1), (-1, -1), "CENTER"),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#BFC7D1")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 4))
    story.append(
        Paragraph(
            "Interpretation: the centralized baseline reaches 98.33% test accuracy. The DP-style runs "
            "show the expected utility cost as noise increases, falling from 97.22% at noise 0.5 to "
            "95.00% at noise 2.0. FedAvg reaches 96.94% under the IID client split, which confirms that "
            "the FL pipeline is functioning before we introduce non-IID data partitions. The MIA AUC "
            "proxy stays close to 0.5, suggesting weak confidence-based membership distinguishability "
            "on this small task; this should be treated as preliminary rather than a complete privacy audit.",
            body,
        )
    )
    story.append(
        KeepTogether(
            [
                Image(str(DP_FIGURE), width=4.65 * inch, height=2.05 * inch),
                Paragraph(
                    "Figure 1. DP noise multiplier versus test accuracy and confidence-based membership-inference proxy.",
                    caption,
                ),
            ]
        )
    )

    story.append(PageBreak())
    story.append(Paragraph("Obstacles and Challenges", h1))
    obstacle_items = [
        "The current dataset is intentionally small, so it validates the pipeline but not final-scale performance.",
        "The DP prototype implements clipping and Gaussian noise, but formal epsilon accounting remains future work.",
        "The FL experiment currently uses IID clients; non-IID splits are likely to create harder convergence behavior.",
        "HE is computationally expensive for training, so the current scope is encrypted inference feasibility.",
    ]
    for item in obstacle_items:
        story.append(Paragraph(f"- {item}", bullet))

    story.append(Paragraph("Engagement with the Literature", h1))
    story.append(
        Paragraph(
            "Abadi et al. introduced DP-SGD as the central mechanism behind our DP prototype: per-example "
            "gradient clipping limits any individual record's contribution, while Gaussian noise provides "
            "the privacy mechanism. Our current implementation captures this operational idea, but the "
            "privacy accountant from the paper is not yet implemented.",
            body,
        )
    )
    story.append(
        Paragraph(
            "McMahan et al. motivates the FL component through FedAvg, where clients train locally and the "
            "server aggregates model updates. Our IID five-client simulation is a first sanity check of this "
            "training loop; the next step is to reproduce the more realistic non-IID setting emphasized by "
            "federated learning literature.",
            body,
        )
    )
    story.append(
        Paragraph(
            "Lee et al. helps frame the HE track. The key lesson is that encrypted neural-network inference "
            "can protect client inputs, but it adds substantial computational overhead and requires careful "
            "model choices. This is why our Part 2 work treats HE as inference-only feasibility rather than "
            "full encrypted training.",
            body,
        )
    )
    story.append(
        Paragraph(
            "Basu et al. reinforces the importance of benchmarking privacy mechanisms under shared tasks "
            "and metrics. We used that idea to keep baseline, DP, and FL results comparable instead of "
            "evaluating each technique in isolation.",
            body,
        )
    )

    story.append(Paragraph("Remaining Plan", h1))
    plan_items = [
        "Add formal DP privacy accounting, preferably using an RDP accountant or Opacus-style epsilon reporting.",
        "Run non-IID FL experiments and compare convergence against the current IID FedAvg curve.",
        "Attempt the optional TenSEAL encrypted linear inference demo and report plaintext versus encrypted latency.",
        "If time allows, scale from digits to MNIST or a small CNN so the final report better matches the original scope.",
        "Prepare final figures, a short demo screenshot, and presentation slides centered on the observed tradeoffs.",
    ]
    for item in plan_items:
        story.append(Paragraph(f"- {item}", bullet))

    story.append(Paragraph("Planned Class Presentation", h1))
    story.append(
        Paragraph(
            "We plan to present the project as a staged comparison: first the non-private baseline, then how "
            "DP changes utility and privacy proxies, then how FL changes the trust and deployment model, and "
            "finally why HE is promising but costly for encrypted inference. The demo will use the Colab "
            "notebook and the generated result table/curves.",
            body,
        )
    )

    story.append(Paragraph("References", h1))
    refs = [
        "Abadi et al., Deep Learning with Differential Privacy, CCS 2016.",
        "McMahan et al., Communication-Efficient Learning of Deep Networks from Decentralized Data, AISTATS 2017.",
        "Lee et al., Privacy-Preserving Machine Learning with Fully Homomorphic Encryption for Deep Neural Network, IEEE Access 2022.",
        "Basu et al., Benchmarking Differential Privacy and Federated Learning for BERT Models, arXiv 2021.",
    ]
    for ref in refs:
        story.append(Paragraph(f"- {ref}", bullet))

    doc.build(story)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_run(run, font_name="Calibri", size=11, bold=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)


def add_docx_heading(doc, text, level=1):
    paragraph = doc.add_heading(level=level)
    run = paragraph.add_run(text)
    style_run(run, size=16 if level == 1 else 13, bold=True, color="2E74B5")
    return paragraph


def add_docx_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    style_run(run, size=11)
    return paragraph


def add_docx_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    style_run(run, size=10.5)
    return paragraph


def build_docx(df):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EEP 595 Privacy-Preserving ML - Project Part 2 Progress Report")
    style_run(run, size=16, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Project: Evaluating Privacy-Utility Tradeoffs in Machine Learning | "
        "Team: Joe Zhang, Mingwei Xu, Houser Zhang | Due: May 16, 2026"
    )
    style_run(run, size=9, color="555555")

    add_docx_heading(doc, "Progress and Execution", 1)
    add_docx_paragraph(
        doc,
        "Our proposal committed to benchmarking differential privacy (DP), federated learning (FL), "
        "and homomorphic encryption (HE) on a shared classification task. For Part 2, we built a "
        "runnable Google Colab pipeline that validates the experimental structure before scaling to "
        "larger models and datasets. The current implementation uses scikit-learn's digits dataset "
        "(1,797 grayscale 8x8 handwritten digit images), normalizes pixels to [0, 1], and uses an "
        "80/20 stratified train-test split with seed 509.",
    )
    add_docx_paragraph(
        doc,
        "Completed components include a non-private centralized logistic-regression baseline, a NumPy "
        "softmax classifier with per-example gradient clipping and Gaussian noise as a DP-SGD-style "
        "prototype, a five-client IID FedAvg simulation, and an HE feasibility track focused on encrypted "
        "inference rather than encrypted training.",
    )

    add_docx_heading(doc, "Preliminary Results", 1)
    table_data = result_table_data(df)
    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
    table.style = "Table Grid"
    widths = [Inches(2.1), Inches(0.8), Inches(0.75), Inches(0.6), Inches(0.7), Inches(0.7)]
    for row_idx, row in enumerate(table_data):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 or row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(text))
            style_run(run, size=8.3, bold=(row_idx == 0), color="0B2545" if row_idx == 0 else None)
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, "E8EEF5")

    add_docx_paragraph(
        doc,
        "The centralized baseline reaches 98.33% test accuracy. The DP-style runs show the expected "
        "utility cost as noise increases, falling from 97.22% at noise 0.5 to 95.00% at noise 2.0. "
        "FedAvg reaches 96.94% under the IID client split, confirming that the FL pipeline works before "
        "we introduce non-IID data partitions. The confidence-based MIA AUC proxy stays close to 0.5, "
        "so it should be treated as a preliminary signal rather than a complete privacy audit.",
    )
    doc.add_picture(str(DP_FIGURE), width=Inches(5.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_docx_heading(doc, "Obstacles and Challenges", 1)
    for item in [
        "The current dataset is intentionally small, so it validates the pipeline but not final-scale performance.",
        "The DP prototype implements clipping and Gaussian noise, but formal epsilon accounting remains future work.",
        "The FL experiment currently uses IID clients; non-IID splits are likely to create harder convergence behavior.",
        "HE is computationally expensive for training, so the current scope is encrypted inference feasibility.",
    ]:
        add_docx_bullet(doc, item)

    add_docx_heading(doc, "Engagement with the Literature", 1)
    add_docx_paragraph(
        doc,
        "Abadi et al. introduced DP-SGD as the central mechanism behind our DP prototype: per-example "
        "gradient clipping limits any individual record's contribution, while Gaussian noise provides the "
        "privacy mechanism. Our current implementation captures this operational idea, but the privacy "
        "accountant from the paper is not yet implemented.",
    )
    add_docx_paragraph(
        doc,
        "McMahan et al. motivates the FL component through FedAvg, where clients train locally and the server "
        "aggregates model updates. Our IID five-client simulation is a first sanity check of this training loop; "
        "the next step is to reproduce the more realistic non-IID setting emphasized by FL literature.",
    )
    add_docx_paragraph(
        doc,
        "Lee et al. helps frame the HE track. Encrypted inference can protect client inputs, but it adds "
        "substantial computational overhead and requires careful model choices. This is why our Part 2 work "
        "treats HE as inference-only feasibility rather than full encrypted training.",
    )
    add_docx_paragraph(
        doc,
        "Basu et al. reinforces the importance of benchmarking privacy mechanisms under shared tasks and "
        "metrics. We used that idea to keep baseline, DP, and FL results comparable instead of evaluating "
        "each technique in isolation.",
    )

    add_docx_heading(doc, "Remaining Plan", 1)
    for item in [
        "Add formal DP privacy accounting, preferably using an RDP accountant or Opacus-style epsilon reporting.",
        "Run non-IID FL experiments and compare convergence against the current IID FedAvg curve.",
        "Attempt the optional TenSEAL encrypted linear inference demo and report plaintext versus encrypted latency.",
        "If time allows, scale from digits to MNIST or a small CNN so the final report better matches the original scope.",
        "Prepare final figures, a short demo screenshot, and presentation slides centered on the observed tradeoffs.",
    ]:
        add_docx_bullet(doc, item)

    add_docx_heading(doc, "Planned Class Presentation", 1)
    add_docx_paragraph(
        doc,
        "We plan to present the project as a staged comparison: first the non-private baseline, then how DP "
        "changes utility and privacy proxies, then how FL changes the trust and deployment model, and finally "
        "why HE is promising but costly for encrypted inference. The demo will use the Colab notebook and the "
        "generated result table/curves.",
    )

    doc.save(DOCX_PATH)


def main():
    df = load_results()
    build_pdf(df)
    build_docx(df)
    print(f"Wrote {PDF_PATH}")
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()
